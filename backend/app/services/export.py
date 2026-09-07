from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


_CJK_FONT = "STSong-Light"
_cjk_font_registered = False


def _ensure_cjk_font() -> None:
    """中文字体只需注册一次，避免重复注册开销。"""
    global _cjk_font_registered
    if not _cjk_font_registered:
        pdfmetrics.registerFont(UnicodeCIDFont(_CJK_FONT))
        _cjk_font_registered = True


def artifact_to_docx(title: str, content: str) -> bytes:
    buffer = BytesIO()
    document = Document()
    document.add_heading(title, level=0)
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#### "):
            document.add_heading(line[5:], level=4)
        elif line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(buffer)
    try:
        return buffer.getvalue()
    finally:
        buffer.close()


def artifact_to_pdf(title: str, content: str) -> bytes:
    _ensure_cjk_font()
    buffer = BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=A4, title=title)
    styles = getSampleStyleSheet()
    for style_name in list(styles.byName):
        styles[style_name].fontName = _CJK_FONT

    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 14)]
    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#### "):
            story.append(Paragraph(escape(line[5:]), styles["Heading4"]))
        elif line.startswith("### "):
            story.append(Paragraph(escape(line[4:]), styles["Heading3"]))
        elif line.startswith("## "):
            story.append(Paragraph(escape(line[3:]), styles["Heading2"]))
        elif line.startswith("# "):
            story.append(Paragraph(escape(line[2:]), styles["Heading1"]))
        elif line.startswith("- ") or line.startswith("* "):
            story.append(Paragraph(f"• {escape(line[2:])}", styles["BodyText"]))
        else:
            story.append(Paragraph(escape(line), styles["BodyText"]))
    pdf.build(story)
    try:
        return buffer.getvalue()
    finally:
        buffer.close()
